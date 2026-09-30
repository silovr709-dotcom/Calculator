#!/usr/bin/env python3
"""Собирает в один JSON ВСЕ справочники/разбивки материалов Висмы из файлов репозитория.

Скрипт не придумывает значения: переносит строки из исходных xlsx/PDF и сохраняет
служебные поля (источник, лист, строка), чтобы любую позицию можно было проверить.

JSON: app/public/data/factory-dicts-visma.json
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from extract_millings import CATALOG as MILLINGS_FILE  # noqa: E402
from extract_millings import load_millings  # noqa: E402

OUT = 'app/public/data/factory-dicts-visma.json'

LDSP_BREAKDOWN_FILES = [
    '!разбивка 2026 ЛДСП.xlsx',
    '!!ЛДСП разбивка.xlsx',
    'ЛДСП.xlsx',
]
FILM_FILES = [
    '!!разбивка ПВХ.xlsx',
    '!!МДФ ПВХ разбивка.xlsx',
]
PLASTIC_FILES = [
    '!разбивка 2026 ARPA, FENIX, AGT, Rexay.xlsx',
    '!!ARPA, FENIX, AGT, Rexay разбивка.xlsx',
    'VM_ARPA, FENIX, AGT 17.03.2025.xlsx',
]
COMPACT_FILE = '!разбивка 2026 компакт Slotex.xlsx'
LAMARTY_THICKNESS_FILE = 'ЛДСП LAMARTY толщина.xlsx'
UVADREV_THICKNESS_FILE = 'ЛДСП УВАДРЕВ толщины.xlsx'

ROOT = Path(__file__).resolve().parents[1]


def clean(text: Any) -> str:
    return re.sub(r'\s+', ' ', str(text)).strip() if text is not None else ''


def norm(text: Any) -> str:
    return clean(text).lower().replace('ё', 'е').replace('×', '*').replace('х', '*')


def source_ref(file: str, sheet: str, row: int) -> dict[str, Any]:
    return {'file': file, 'sheet': sheet, 'row': row}


def add_source(item: dict[str, Any], file: str, sheet: str, row: int) -> None:
    ref = source_ref(file, sheet, row)
    sources = item.setdefault('sources', [])
    if ref not in sources:
        sources.append(ref)


def merge_items(items: list[dict[str, Any]], key_fields: list[str], prefer_first: bool = True) -> list[dict[str, Any]]:
    """Склеивает одинаковые позиции, но не теряет источники и непустые поля."""
    merged: dict[tuple[str, ...], dict[str, Any]] = {}
    order: list[tuple[str, ...]] = []
    for item in items:
        key = tuple(norm(item.get(field)) for field in key_fields)
        if not any(key):
            key = tuple([norm(item.get('name')), norm(item.get('article') or item.get('code'))])
        current = merged.get(key)
        if current is None:
            merged[key] = item
            order.append(key)
            continue
        for source in item.get('sources', []):
            if source not in current.setdefault('sources', []):
                current['sources'].append(source)
        for k, v in item.items():
            if k == 'sources':
                continue
            if v in (None, '', [], {}):
                continue
            if current.get(k) in (None, '', [], {}):
                current[k] = v
            elif not prefer_first and current.get(k) != v:
                current[f'{k}Alt'] = v
    return [merged[key] for key in order]


def article_from_text(text: str) -> str | None:
    # Увадрев/ЛДСП: U1655, U31148, W1000, H1180 и т.п.
    m = re.search(r'\b([A-ZА-Я]{1,3}\d{3,6}(?:/[A-ZА-ЯМM])?)\b', text)
    if m:
        return m.group(1)
    # Пластики/декоры часто выглядят как 0029, 1456/Papier, AGT4560/383.
    m = re.search(r'\b([A-Za-zА-Яа-я]*\d{2,6}(?:/[\wА-Яа-я]+)?)\b', text)
    return m.group(1) if m else None


# --------------------------------------------------------------------------- ЛДСП

def extract_ldsp_breakdowns() -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for file in LDSP_BREAKDOWN_FILES:
        path = ROOT / file
        if not path.exists():
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb['ЛДСП']
        brand = ''
        fmt = ''
        category = ''
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            c0 = clean(row[0]) if row else ''
            if not c0:
                continue
            lowered = c0.lower()
            if c0.startswith('!!!') or lowered.startswith('лдсп') or 'артикул' in lowered or 'кромка' in lowered:
                continue
            brand_match = re.match(r'^(.+?)\s+(\d{3,4}\s*[*xх×]\s*\d{3,4}\s*[*xх×]\s*\d{1,2})(?:\s|$)', c0, flags=re.IGNORECASE)
            if brand_match:
                brand = clean(brand_match.group(1))
                fmt = clean(brand_match.group(2)).replace('х', '*').replace('×', '*').replace(' ', '')
                category = ''
                continue
            cat_match = re.match(r'^(\d)\s*категор', lowered)
            if cat_match:
                category = f'{cat_match.group(1)} категория'
                continue
            if not brand:
                continue
            if len(c0) < 3 or lowered in {'база', 'цвет'}:
                continue
            texture = c0.startswith('!')
            name = c0.lstrip('!').strip()
            item = {
                'name': name,
                'article': article_from_text(name),
                'texture': texture,
                'brand': brand,
                'format': fmt,
                'category': category,
                'edgingArticle': clean(row[6]) or None,
                'manufacturerEdge04': clean(row[7]) or None,
                'manufacturerEdge1or2': clean(row[8]) or None,
                'sheetSurcharge': clean(row[9]) or None,
                'overuseRule': clean(row[10]) or None,
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)
        wb.close()
    return merge_items(items, ['brand', 'article', 'name', 'format'], prefer_first=True)


def extract_lamarty_thickness() -> list[dict[str, Any]]:
    file = LAMARTY_THICKNESS_FILE
    path = ROOT / file
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['складская программа']
    headers = [clean(v) for v in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    items: list[dict[str, Any]] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        name = clean(row[1]) if len(row) > 1 else ''
        if not name:
            continue
        thicknesses = {headers[col]: clean(row[col]) for col in range(3, min(9, len(row))) if headers[col] and clean(row[col])}
        if not thicknesses:
            continue
        item = {
            'brand': 'Lamarty',
            'name': name,
            'textureCode': clean(row[2]) or None,
            'thicknesses': thicknesses,
        }
        add_source(item, file, ws.title, row_idx)
        items.append(item)
    wb.close()
    return items


def extract_uvadrev_thickness_and_edges() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    file = UVADREV_THICKNESS_FILE
    path = ROOT / file
    if not path.exists():
        return [], []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    stock: list[dict[str, Any]] = []
    ws = wb['Склад ']
    headers = [clean(v) for v in next(ws.iter_rows(min_row=4, max_row=4, values_only=True))]
    category = ''
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        name = clean(row[2]) if len(row) > 2 else ''
        if not name:
            continue
        category = clean(row[1]) or category
        thicknesses = {headers[col]: clean(row[col]) for col in range(4, min(9, len(row))) if headers[col] and clean(row[col])}
        edge_availability = {headers[col]: clean(row[col]) for col in range(9, min(len(headers), len(row))) if headers[col] and clean(row[col])}
        item = {
            'brand': 'Увадрев',
            'category': category,
            'name': name,
            'article': article_from_text(name),
            'textureCode': clean(row[3]) or None,
            'thicknesses': thicknesses,
            'edgeAvailability': edge_availability,
        }
        add_source(item, file, ws.title, row_idx)
        stock.append(item)

    edges: list[dict[str, Any]] = []
    ws = wb['Подбор ']
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = clean(row[0]) if row else ''
        article = clean(row[1]) if len(row) > 1 else ''
        edge = clean(row[2]) if len(row) > 2 else ''
        if not name or not edge:
            continue
        item = {'brand': 'Увадрев', 'name': name, 'article': article or article_from_text(name), 'edge': edge}
        add_source(item, file, ws.title, row_idx)
        edges.append(item)
    wb.close()
    return stock, edges


# --------------------------------------------------------------------------- ПВХ плёнки
FILM_STATUS = [('выведен', 'выведена'), ('снят', 'снята')]


def parse_film_cell(code_raw: Any, name_raw: Any) -> dict[str, Any] | None:
    code = clean(code_raw)
    name = clean(name_raw)
    if not name or 'пленки нет' in name.lower():
        return None
    lowered_code = code.lower()
    flags: dict[str, Any] = {'texture': '!' in code}
    if code.startswith('**') or lowered_code.startswith('**'):
        flags['onlyMdf16'] = True
    elif code.startswith('*') or lowered_code.startswith('*'):
        flags['onlyMillingMilo'] = True
    codes = code.lstrip('*!').strip()
    codes = re.sub(r'\s+', ' ', codes)
    if not codes or codes.endswith(')') or 'разбивка' in codes.lower():
        return None
    status = 'в работе'
    lowered = name.lower()
    for marker, value in FILM_STATUS:
        if marker in lowered:
            status = value
            break
    for marker, _ in FILM_STATUS:
        name = re.sub(rf'\s*{marker}\w*\s*', ' ', name, flags=re.IGNORECASE).strip()
    return {'code': codes, 'name': name, 'status': status, **flags}


def extract_films() -> list[dict[str, Any]]:
    all_items: list[dict[str, Any]] = []
    block = re.compile(r'^(.+?)\s*\(\s*(\d)\s*кат\s*\)$', flags=re.IGNORECASE)
    for file in FILM_FILES:
        path = ROOT / file
        if not path.exists():
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb['Лист1']
        brand, category = '', ''
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            c0 = clean(row[0]) if row else ''
            m = block.match(c0)
            if m and not any(ch.isdigit() for ch in m.group(1)):
                brand, category = clean(m.group(1)), f'{m.group(2)} кат'
                continue
            for off in (0, 6, 12, 18):
                code_raw = row[off] if len(row) > off else None
                name_raw = row[off + 2] if len(row) > off + 2 else None
                if not code_raw or not name_raw:
                    continue
                item = parse_film_cell(code_raw, name_raw)
                if item and brand:
                    item.update({'brand': brand, 'category': category})
                    add_source(item, file, ws.title, row_idx)
                    all_items.append(item)
        wb.close()
    return merge_items(all_items, ['brand', 'code', 'name'], prefer_first=True)


# --------------------------------------------------------------------------- Пластики/АГТ/FENIX/Rexay/ABET
CATEGORY_RE = re.compile(r'^(\d\s*(?:к|кат|категория)|снят|вывед)', flags=re.IGNORECASE)
DIM_RE = re.compile(r'\d{3,4}\s*[*xх×]\s*\d{3,4}(?:\s*[*xх×]\s*\d[\d,.]*)?', flags=re.IGNORECASE)


def status_from(*values: Any) -> tuple[str, str | None]:
    text = ' '.join(clean(v).lower() for v in values)
    if 'снят' in text:
        return 'снята', None
    if 'вывед' in text:
        return 'выведена', None
    return 'в работе', None


def normalize_category(value: str) -> str | None:
    v = clean(value)
    if not v:
        return None
    if 'снят' in v.lower() or 'вывед' in v.lower():
        return None
    return v


def first_dimension(text: str) -> str | None:
    m = DIM_RE.search(text.replace(',', '.'))
    if not m:
        return None
    return m.group(0).replace('х', '*').replace('×', '*').replace(' ', '')


def parse_article_name(raw: str) -> tuple[str, str]:
    text = clean(raw).lstrip('!').strip()
    parts = text.split(maxsplit=1)
    if not parts:
        return '', ''
    if len(parts) == 1:
        return parts[0], ''
    return parts[0], parts[1]


def extract_plastics_from_file(file: str) -> list[dict[str, Any]]:
    path = ROOT / file
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items: list[dict[str, Any]] = []

    # ARPA: классическая 4-колоночная разбивка.
    if 'ARPA' in wb.sheetnames:
        ws = wb['ARPA']
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            article = clean(row[0]) if row else ''
            fmt = clean(row[1]) if len(row) > 1 else ''
            name = clean(row[2]) if len(row) > 2 else ''
            category_raw = clean(row[3]) if len(row) > 3 else ''
            if not article or not name:
                continue
            status, _ = status_from(category_raw)
            item = {
                'article': article.lstrip('!').strip(),
                'name': name,
                'brand': 'ARPA',
                'format': fmt or first_dimension(name),
                'category': normalize_category(category_raw),
                'status': status,
                'texture': article.startswith('!'),
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)

    # ABET: коллекции в первом столбце, затем: артикул+название / параметры / категория.
    if 'ABET' in wb.sheetnames:
        ws = wb['ABET']
        collection = ''
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            c0 = clean(row[0]) if row else ''
            c1 = clean(row[1]) if len(row) > 1 else ''
            c2 = clean(row[2]) if len(row) > 2 else ''
            if not c0:
                continue
            if not c1 and not CATEGORY_RE.match(c2):
                collection = c0.strip()
                continue
            if not CATEGORY_RE.match(c2) and not c1:
                continue
            article, name = parse_article_name(c0)
            if not article or not name:
                continue
            status, _ = status_from(c2, c1)
            item = {
                'article': article,
                'name': name,
                'brand': 'ABET',
                'collection': collection or None,
                'format': first_dimension(c1),
                'category': normalize_category(c2),
                'status': status,
                'texture': c0.startswith('!') or c1.strip().startswith('*') or '*, ' in c1,
                'details': c1 or None,
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)

    # АГТ: широкий лист; нужны артикул, наименование, категория, артикул кромки.
    if 'АГТ' in wb.sheetnames:
        ws = wb['АГТ']
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            article_raw = clean(row[0]) if row else ''
            name = clean(row[1]) if len(row) > 1 else ''
            if not article_raw or not name or article_raw.lower().startswith('артикул'):
                continue
            if 'категория' in article_raw.lower() and not name:
                continue
            category_raw = clean(row[8]) if len(row) > 8 else ''
            edge = clean(row[9]) if len(row) > 9 else ''
            status, _ = status_from(category_raw)
            item = {
                'article': article_raw.lstrip('!').strip(),
                'name': name,
                'brand': 'AGT',
                'format': first_dimension(name),
                'category': normalize_category(category_raw),
                'status': status,
                'texture': article_raw.startswith('!'),
                'edge': edge or None,
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)

    # Rexay: коллекции отдельными строками, категория может стоять только у первой строки блока.
    if 'Rexay' in wb.sheetnames:
        ws = wb['Rexay']
        collection = ''
        category = ''
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            c0 = clean(row[0]) if row else ''
            fmt = clean(row[1]) if len(row) > 1 else ''
            cat = clean(row[2]) if len(row) > 2 else ''
            if not c0 or c0.lower().startswith('артикул'):
                continue
            if not fmt:
                collection = c0
                category = ''
                continue
            if cat:
                category = cat
            article, name = parse_article_name(c0)
            if not article:
                continue
            status, _ = status_from(cat, c0)
            item = {
                'article': article,
                'name': name,
                'brand': 'Rexay',
                'collection': collection or None,
                'format': fmt or first_dimension(c0),
                'category': normalize_category(category),
                'status': status,
                'texture': c0.startswith('!'),
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)

    # FENIX: артикул / имя / размер / кромка / категория.
    if 'FENIX' in wb.sheetnames:
        ws = wb['FENIX']
        collection = ''
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            article_raw = clean(row[0]) if row else ''
            name = clean(row[1]) if len(row) > 1 else ''
            if not article_raw or article_raw.lower().startswith('артикул'):
                continue
            if not name:
                collection = article_raw
                continue
            fmt = clean(row[2]) if len(row) > 2 else ''
            edge = clean(row[6]) if len(row) > 6 else ''
            category_raw = clean(row[7]) if len(row) > 7 else ''
            status, _ = status_from(category_raw)
            item = {
                'article': article_raw.lstrip('!').strip(),
                'name': name,
                'brand': 'FENIX',
                'collection': collection or None,
                'format': fmt or first_dimension(name),
                'category': normalize_category(category_raw),
                'status': status,
                'texture': article_raw.startswith('!'),
                'edge': edge or None,
            }
            add_source(item, file, ws.title, row_idx)
            items.append(item)

    wb.close()
    return items


def extract_plastics() -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for file in PLASTIC_FILES:
        items.extend(extract_plastics_from_file(file))
    return merge_items(items, ['brand', 'article', 'name', 'format'], prefer_first=True)


# --------------------------------------------------------------------------- Компакт Slotex

def extract_compact_slotex() -> tuple[list[dict[str, Any]], dict[str, str]]:
    file = COMPACT_FILE
    path = ROOT / file
    if not path.exists():
        return [], {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Кратность реализации HPLCompact']

    legend: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=9, values_only=True):
        key = clean(row[1]) if len(row) > 1 else ''
        text = clean(row[2]) if len(row) > 2 else ''
        if key and text:
            legend[key] = text

    product_labels: dict[int, str] = {}
    product = ''
    length = ''
    row_product = next(ws.iter_rows(min_row=11, max_row=11, values_only=True))
    row_length = next(ws.iter_rows(min_row=12, max_row=12, values_only=True))
    row_width = next(ws.iter_rows(min_row=13, max_row=13, values_only=True))
    row_thick = next(ws.iter_rows(min_row=14, max_row=14, values_only=True))
    for col in range(6, 14):
        product = clean(row_product[col]) or product
        length = clean(row_length[col]) or length
        width = clean(row_width[col])
        thick = clean(row_thick[col])
        if product and length and width and thick:
            product_labels[col] = f'{product} {length}×{width}×{thick}'

    items: list[dict[str, Any]] = []
    collection = ''
    for row_idx, row in enumerate(ws.iter_rows(min_row=15, values_only=True), start=15):
        series = clean(row[1]) if len(row) > 1 else ''
        code = clean(row[2]) if len(row) > 2 else ''
        if series and not code:
            collection = series
            continue
        if not code:
            continue
        availability = {label: clean(row[col]) for col, label in product_labels.items() if clean(row[col])}
        item = {
            'brand': 'Slotex',
            'collection': collection or series or None,
            'series': series or None,
            'code': code,
            'name': clean(row[5]) if len(row) > 5 else '',
            'textureCode': clean(row[3]) or None,
            'surfaceType': clean(row[4]) or None,
            'availability': availability,
        }
        add_source(item, file, ws.title, row_idx)
        items.append(item)
    wb.close()
    return items, legend


# --------------------------------------------------------------------------- main

def main() -> None:
    millings, milling_issues = load_millings()
    uvadrev_stock, uvadrev_edges = extract_uvadrev_thickness_and_edges()
    compact_items, compact_legend = extract_compact_slotex()
    sources = [
        *LDSP_BREAKDOWN_FILES,
        *FILM_FILES,
        *PLASTIC_FILES,
        COMPACT_FILE,
        LAMARTY_THICKNESS_FILE,
        UVADREV_THICKNESS_FILE,
        MILLINGS_FILE,
    ]
    data = {
        'id': 'visma-dicts-2026-complete',
        'name': 'Справочники разбивок Висма: цвета, материалы, кромки и кратность',
        'sources': sources,
        'groups': {
            'ldspColors': {
                'label': 'Цвета ЛДСП: разбивки, кромки, доплаты, перерасход',
                'legend': {
                    'texture': '«!» — направление рисунка/текстура: обязательно указывать в бланке',
                    'edgingArticle': 'Кромка GP (база) 0,4×19 без доплаты',
                    'manufacturerEdge04': 'Кромка производителя 0,4×19 за доплату',
                    'manufacturerEdge1or2': 'Кромка производителя 1/2 мм за доплату по правилам листа',
                    'sheetSurcharge': 'Доплата за лист из разбивки',
                    'overuseRule': 'Перерасход м² досчитывается кратно листу по разбивке',
                },
                'items': extract_ldsp_breakdowns(),
            },
            'ldspThickness': {
                'label': 'ЛДСП: складские толщины и наличие кромок',
                'legend': {
                    'thicknesses': '● — доступная толщина/складская программа по исходному листу',
                    'edgeAvailability': 'MB/○/● — обозначения из листа Увадрев, перенесены без интерпретации',
                },
                'items': [*extract_lamarty_thickness(), *uvadrev_stock],
            },
            'ldspEdges': {
                'label': 'ЛДСП Увадрев: подбор кромки к декору',
                'items': uvadrev_edges,
            },
            'films': {
                'label': 'Плёнки для фасадов МДФ (ПВХ): все категории и статусы',
                'legend': {
                    'texture': '«!» — направление рисунка: обязательно указывать в бланке',
                    'onlyMillingMilo': '«*» — только с фрезеровкой «Мыло»',
                    'onlyMdf16': '«**» — только фрезеровки на МДФ 16 мм',
                    'status': '«выведена»/«снята» — в работу не берутся',
                },
                'items': extract_films(),
            },
            'plastics': {
                'label': 'Пластики/HPL/AGT/FENIX/Rexay/ABET: разбивки и кромки',
                'legend': {
                    'texture': '«!» или * в источнике — направление/рисунок; знак сохранён как признак текстуры',
                    'status': '«снят»/«выведен» из категории/примечаний — позиция не берётся в работу',
                    'edge': 'Кромка/артикул кромки из соответствующего листа',
                },
                'items': extract_plastics(),
            },
            'compactHpl': {
                'label': 'Компакт Slotex: кратность реализации HPL Compact / SolidTop',
                'legend': compact_legend,
                'items': compact_items,
            },
            'millings': {
                'label': 'Фрезеровки фасадов (каталог 2026)',
                'legend': {
                    'category': 'Категория 1–4 — ею прайс 2026 задаёт цену фасада (ПВХ и эмаль)',
                    'coatings': '«Возможность изготовления» — в чём фабрика делает фрезеровку',
                    'sizes': 'Размеры — допустимые «высота*ширина» в мм для глухого, РК, РБК и ящика',
                    'source': 'Источник — страница PDF каталога, по ней значение можно перепроверить',
                },
                'issues': milling_issues,
                'items': millings,
            },
        },
    }
    total = sum(len(g['items']) for g in data['groups'].values())
    with open(ROOT / OUT, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    for key, group in data['groups'].items():
        print(f"{key}: {len(group['items'])}")
    print(f'ИТОГО {total} позиций → {OUT}')
    if total < 1000:
        sys.exit('слишком мало позиций — проверить парсер')


if __name__ == '__main__':
    main()
